import streamlit as st
import requests
import io
import pandas as pd
import time
from openai import OpenAI
import anthropic
import google.generativeai as genai
from fpdf import FPDF
from PyPDF2 import PdfReader
from docx import Document
from scraper import get_video_urls, extract_video_id, get_transcript
from ui import apply_css, render_header, render_footer
from algodev import simulate_dfil, simulate_gating, generate_seo_report

apply_css()
render_header()

tab1, tab2 = st.tabs(["Extract","Query"])

with tab1:
    url_input = st.text_input("YouTube URL")
    max_videos = st.text_input("Max videos (channel)")
    uploaded_file = st.file_uploader("PDF/DOC", type=["pdf","docx"])

    if "videos" not in st.session_state:
        st.session_state.videos = None
    if "selected_videos" not in st.session_state:
        st.session_state.selected_videos = []

    if st.button("Load"):
        full_text = ""
        report_data = []
        if uploaded_file:
            with st.spinner("Document"):
                doc_text = ""
                try:
                    if uploaded_file.type == "application/pdf":
                        reader = PdfReader(uploaded_file)
                        doc_text = "".join(page.extract_text() or "" for page in reader.pages)
                    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                        doc = Document(uploaded_file)
                        doc_text = "\n".join(p.text for p in doc.paragraphs)
                    full_text += "\n\n--- Document ---\n" + doc_text
                    report_data.append({'title':uploaded_file.name,'transcript':doc_text})
                    st.success("Document ready")
                except Exception as e:
                    st.error(f"Document failed: {e}")

        if url_input:
            with st.spinner("Fetching"):
                videos, error = get_video_urls(url_input.strip(), int(max_videos) if max_videos.isdigit() else None)
                if error:
                    st.error(error)
                else:
                    st.session_state.videos = videos
                    st.session_state.selected_videos = videos  # default all selected

        st.success("Loaded")

    if st.session_state.videos:
        st.markdown("Select videos")
        cols = st.columns(4)
        selected_videos = []
        for idx, v in enumerate(st.session_state.videos):
            with cols[idx % 4]:
                st.image(v['thumbnail'], width=200)
                checked = st.checkbox(v['title'][:40] + "..." if len(v['title']) > 40 else v['title'], 
                                     value=v in st.session_state.selected_videos, 
                                     key=f"chk_{v['video_id']}")
                if checked:
                    selected_videos.append(v)
        st.session_state.selected_videos = selected_videos

    if st.session_state.selected_videos or uploaded_file:
        if st.button("Extract"):
            start_time = time.time()
            progress_bar = st.progress(0)
            status_text = st.empty()
            eta_text = st.empty()
            total_time_text = st.empty()

            total = len(st.session_state.selected_videos)
            times = []

            full_text = ""
            report_data = []

            for i, v in enumerate(st.session_state.selected_videos):
                loop_start = time.time()
                status_text.text(f"({i+1}/{total}) {v['title']}")
                trans = get_transcript(v['url'])
                if trans:
                    report_data.append({**v, 'transcript': trans})
                    full_text += f"\n\n--- {v['title']} ---\n{trans}"
                progress_bar.progress((i+1) / total)
                times.append(time.time() - loop_start)
                if times:
                    avg = sum(times) / len(times)
                    rem_sec = (total - i - 1) * avg
                    rem_min = int(rem_sec // 60)
                    rem_sec = int(rem_sec % 60)
                    eta_text.text(f"Remaining: {rem_min} min {rem_sec:02d} sec")

            total_sec = time.time() - start_time
            total_min = int(total_sec // 60)
            total_sec = int(total_sec % 60)
            total_time_text.text(f"Total: {total_min} min {total_sec:02d} sec")
            status_text.text("Done")

            if full_text:
                st.session_state.full_text = full_text
                st.session_state.report_data = report_data

                buf = io.BytesIO()
                pdf = FPDF()
                pdf.set_auto_page_break(auto=True, margin=15)
                for item in report_data:
                    pdf.add_page()
                    pdf.set_font("Helvetica", 'B', 16)
                    safe_title = item['title'].encode('latin-1', 'ignore').decode('latin-1')
                    pdf.multi_cell(0, 10, safe_title, align='C')
                    pdf.ln(5)
                    pdf.set_font("Helvetica", '', 10)
                    pdf.cell(0, 8, f"Source: {'YouTube' if 'url' in item else 'Document'}", ln=True)
                    pdf.ln(10)
                    pdf.set_font("Helvetica", '', 11)
                    safe_transcript = item['transcript'].encode('latin-1', 'ignore').decode('latin-1')
                    pdf.multi_cell(0, 6, safe_transcript)
                pdf.output(buf)
                buf.seek(0)
                st.download_button("Download PDF", buf, "Content.pdf", "application/pdf")

                with st.spinner("AlgoDev"):
                    ranked_terms, G, all_terms, net, extra_msg = simulate_dfil(full_text)
                    context_mode = st.selectbox("Context", ["Evergreen", "QDF", "Mobile"])
                    gated_weights = simulate_gating(ranked_terms, G, context_mode, st.session_state.get('serper_key'), url_input or (uploaded_file.name if uploaded_file else ""))
                    report, csv_data = generate_seo_report(ranked_terms, G, gated_weights, extra_msg)

                    st.text_area("Report", report, height=400)

                    if csv_data is not None and not csv_data.empty:
                        csv_buffer = io.StringIO()
                        csv_data.to_csv(csv_buffer, index=False)
                        st.download_button("Download CSV", csv_buffer.getvalue(), "keywords.csv", "text/csv")

                    if net is not None:
                        st.markdown("Graph")
                        net.save_graph("graph.html")
                        with open("graph.html", "r", encoding="utf-8") as f:
                            st.components.v1.html(f.read(), height=650, scrolling=True)
                    else:
                        st.info("No graph")

                st.success("Done — go to Query")

with tab2:
    if 'full_text' not in st.session_state:
        st.info("Extract first")
    else:
        st.markdown("Query")
        if 'messages' not in st.session_state:
            st.session_state.messages = []
        for msg in st.session_state.messages:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])
        query = st.chat_input("Ask...")
        if query:
            st.session_state.messages.append({"role": "user", "content": query})
            with st.chat_message("user"):
                st.markdown(query)
            context = st.session_state.full_text[:18000]
            web_context = ""
            if st.session_state.get('serper_key'):
                with st.spinner("Web"):
                    try:
                        payload = {"q": query}
                        headers = {"X-API-KEY": st.session_state.serper_key, "Content-Type": "application/json"}
                        resp = requests.post("https://google.serper.dev/search", json=payload, headers=headers)
                        if resp.status_code == 200:
                            results = resp.json().get('organic', [])[:4]
                            web_context = "\n".join([f"- {r['title']}: {r['snippet']}" for r in results])
                    except:
                        web_context = "(Web failed)"
            base_prompt = f"""
Use only this content + web.
Analyze critically, teach deeply, show flaws/risks/alternatives.
Content:
{context}
Web:
{web_context}
Question: {query}
"""
            responses = {}
            with st.spinner("Querying"):
                if st.session_state.get('openai_key'):
                    try:
                        client = OpenAI(api_key=st.session_state.openai_key)
                        resp = client.chat.completions.create(model="gpt-4o-mini", messages=[{"role": "user", "content": base_prompt}])
                        responses["OpenAI"] = resp.choices[0].message.content
                    except Exception as e:
                        responses["OpenAI"] = f"Error: {e}"
                if st.session_state.get('anthropic_key'):
                    try:
                        client = anthropic.Anthropic(api_key=st.session_state.anthropic_key)
                        resp = client.messages.create(model="claude-3-5-sonnet-20241022", max_tokens=2048, messages=[{"role": "user", "content": base_prompt}])
                        responses["Claude"] = resp.content[0].text
                    except Exception as e:
                        responses["Claude"] = f"Error: {e}"
                if st.session_state.get('gemini_key'):
                    try:
                        genai.configure(api_key=st.session_state.gemini_key)
                        model = genai.GenerativeModel('gemini-1.5-pro')
                        resp = model.generate_content(base_prompt)
                        responses["Gemini"] = resp.text
                    except Exception as e:
                        responses["Gemini"] = f"Error: {e}"
                if st.session_state.get('grok_key'):
                    try:
                        client = OpenAI(api_key=st.session_state.grok_key, base_url="https://api.x.ai/v1")
                        resp = client.chat.completions.create(model="grok-beta", messages=[{"role": "user", "content": base_prompt}])
                        responses["Grok"] = resp.choices[0].message.content
                    except Exception as e:
                        responses["Grok"] = f"Error: {e}"
            for model, text in responses.items():
                with st.expander(model, expanded=True):
                    st.markdown(text)
            combined = "\n\n---\n\n".join([f"**{m}:**\n{t}" for m, t in responses.items()])
            st.session_state.messages.append({"role": "assistant", "content": combined})

render_footer()
